"""API endpoints for the importer app."""

from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.urls import include, path

from drf_spectacular.utils import extend_schema
from rest_framework import serializers
from rest_framework.exceptions import PermissionDenied
from rest_framework.response import Response
from rest_framework.views import APIView

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

import importer.models
import importer.registry
import importer.serializers
import InvenTree.permissions
from InvenTree.api import BulkDeleteMixin
from InvenTree.filters import SEARCH_ORDER_FILTER
from InvenTree.mixins import (
    CreateAPI,
    ListAPI,
    ListCreateAPI,
    RetrieveUpdateAPI,
    RetrieveUpdateDestroyAPI,
)
from users.permissions import check_user_permission


class DataImporterPermissionMixin:
    """Mixin class for checking permissions on DataImporter objects."""

    # Default permissions: User must be authenticated
    permission_classes = [
        InvenTree.permissions.IsAuthenticatedOrReadScope,
        InvenTree.permissions.DataImporterPermission,
    ]


class DataImporterModelSerializer(serializers.Serializer):
    """Model references to map info that might get imported."""

    serializer = serializers.CharField(read_only=True)
    model_type = serializers.CharField(read_only=True)
    api_url = serializers.URLField(read_only=True, allow_null=True)


class DataImporterModelList(APIView):
    """API endpoint for displaying a list of models available for import."""

    permission_classes = [InvenTree.permissions.IsAuthenticatedOrReadScope]
    serializer_class = DataImporterModelSerializer(many=True)

    def get(self, request):
        """Return a list of models available for import."""
        models = []

        for serializer in importer.registry.get_supported_serializers():
            model = serializer.Meta.model
            url = model.get_api_url() if hasattr(model, 'get_api_url') else None

            models.append({
                'serializer': str(serializer.__name__),
                'model_type': model.__name__.lower(),
                'api_url': url,
            })

        return Response(models)


class DataImportSessionMixin:
    """Mixin class for DataImportSession API views."""

    queryset = importer.models.DataImportSession.objects.all()
    serializer_class = importer.serializers.DataImportSessionSerializer
    permission_classes = [InvenTree.permissions.DataImporterPermission]

    def get_queryset(self):
        """Return the set of DataImportSession objects that the user has permission to view."""
        queryset = super().get_queryset()

        try:
            user = self.request.user
        except AttributeError:
            raise PermissionDenied('User information is not available')

        # Allow staff users access to all DataImportSession objects
        if user.is_staff:
            return queryset

        # For non-staff users, only allow access to sessions that they have created
        return queryset.filter(user=user)


class DataImportSessionList(BulkDeleteMixin, DataImportSessionMixin, ListCreateAPI):
    """API endpoint for accessing a list of DataImportSession objects."""

    filter_backends = SEARCH_ORDER_FILTER
    filterset_fields = ['model_type', 'status', 'user']
    ordering_fields = ['timestamp', 'status', 'model_type']


class DataImportSessionDetail(DataImportSessionMixin, RetrieveUpdateDestroyAPI):
    """Detail endpoint for a single DataImportSession object."""


class DataImportSessionAcceptFields(APIView):
    """API endpoint to accept the field mapping for a DataImportSession."""

    permission_classes = [InvenTree.permissions.IsAuthenticatedOrReadScope]
    serializer_class = None

    @extend_schema(
        responses={200: importer.serializers.DataImportSessionSerializer(many=False)}
    )
    def post(self, request, pk):
        """Accept the field mapping for a DataImportSession."""
        session = get_object_or_404(importer.models.DataImportSession, pk=pk)

        # Check that the user has permission to accept the field mapping
        if model_class := session.model_class:
            if not check_user_permission(request.user, model_class, 'change'):
                raise PermissionDenied()

        # Attempt to accept the mapping (may raise an exception if the mapping is invalid)
        session.accept_mapping()

        return Response(importer.serializers.DataImportSessionSerializer(session).data)


class DataImportSessionAcceptRows(DataImporterPermissionMixin, CreateAPI):
    """API endpoint to accept the rows for a DataImportSession."""

    queryset = importer.models.DataImportSession.objects.all()
    serializer_class = importer.serializers.DataImportAcceptRowSerializer

    def get_serializer_context(self):
        """Add the import session object to the serializer context."""
        ctx = super().get_serializer_context()

        try:
            ctx['session'] = importer.models.DataImportSession.objects.get(
                pk=self.kwargs.get('pk', None)
            )
        except Exception:
            pass

        ctx['request'] = self.request
        return ctx


class DataImportColumnMappingList(DataImporterPermissionMixin, ListAPI):
    """API endpoint for accessing a list of DataImportColumnMap objects."""

    queryset = importer.models.DataImportColumnMap.objects.all()
    serializer_class = importer.serializers.DataImportColumnMapSerializer

    filter_backends = SEARCH_ORDER_FILTER

    filterset_fields = ['session']


class DataImportColumnMappingDetail(DataImporterPermissionMixin, RetrieveUpdateAPI):
    """Detail endpoint for a single DataImportColumnMap object."""

    queryset = importer.models.DataImportColumnMap.objects.all()
    serializer_class = importer.serializers.DataImportColumnMapSerializer


class DataImportRowList(DataImporterPermissionMixin, BulkDeleteMixin, ListAPI):
    """API endpoint for accessing a list of DataImportRow objects."""

    queryset = importer.models.DataImportRow.objects.all()
    serializer_class = importer.serializers.DataImportRowSerializer

    filter_backends = SEARCH_ORDER_FILTER

    filterset_fields = ['session', 'valid', 'complete']

    ordering_fields = ['pk', 'row_index', 'valid']

    ordering = 'row_index'


class DataImportRowDetail(DataImporterPermissionMixin, RetrieveUpdateDestroyAPI):
    """Detail endpoint for a single DataImportRow object."""

    queryset = importer.models.DataImportRow.objects.all()
    serializer_class = importer.serializers.DataImportRowSerializer


# Template field definitions for each model type
IMPORT_TEMPLATES = {
    'part': {
        'label': '物料导入模板',
        'fields': [
            ('name', '物料名称', '物料名称'),
            ('IPN', '物料型号', '物料型号'),
            ('description', '描述', '描述'),
            ('category', '分类ID或路径', '例如 5 或 CatA/CatB'),
            ('keywords', '关键词', '逗号分隔'),
            ('units', '单位', '个/件/m/kg'),
            ('active', '启用', 'True/False'),
            ('assembly', '是装配件', 'True/False'),
            ('component', '是组件', 'True/False'),
            ('purchaseable', '可采购', 'True/False'),
            ('salable', '可销售', 'True/False'),
            ('trackable', '可跟踪', 'True/False'),
            ('virtual', '虚拟件', 'True/False'),
            ('is_template', '是模板', 'True/False'),
            ('variant_of', '父模板ID', ''),
            ('revision', '版本号', ''),
            ('link', '链接', ''),
            ('minimum_stock', '最低库存', '数字'),
            ('maximum_stock', '最高库存', '数字'),
            ('default_expiry', '默认有效期(天)', '数字'),
        ],
    },
    'partcategory': {
        'label': '物料分类导入模板',
        'fields': [
            ('name', '分类名称', '必填'),
            ('description', '描述', ''),
            ('parent', '父分类ID', ''),
            ('icon', '图标', ''),
            ('structural', '结构分类', 'True/False'),
        ],
    },
    'stockitem': {
        'label': '库存导入模板',
        'fields': [
            ('part', '物料ID', '必填'),
            ('location', '库位ID', ''),
            ('quantity', '数量', '必填'),
            ('serial', '序列号', ''),
            ('batch', '批次号', ''),
            ('status', '状态', ''),
            ('notes', '备注', ''),
        ],
    },
}


class DataImportTemplateView(APIView):
    """API endpoint to download an import template file (XLSX)."""

    permission_classes = [InvenTree.permissions.IsAuthenticatedOrReadScope]

    @extend_schema(
        parameters=[
            serializers.Serializer('model', serializers.CharField()),
        ],
        responses={200: None},
    )
    def get(self, request):
        model_type = request.query_params.get('model', 'part').lower()

        if model_type not in IMPORT_TEMPLATES:
            return Response(
                {'error': f'Unsupported model type: {model_type}'},
                status=400,
            )

        template = IMPORT_TEMPLATES[model_type]

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = template['label']

        # ── Styles ──
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin'),
        )
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_font = Font(name='微软雅黑', bold=True, color='FFFFFF', size=11)
        header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

        hint_fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
        hint_font = Font(name='微软雅黑', size=9, color='888888')
        hint_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

        data_font = Font(name='微软雅黑', size=10)
        data_align = Alignment(vertical='center', wrap_text=True)

        # ── Build template: Row 1 = hints, Row 2 = field names, Row 3+ = data ──
        fields = template['fields']
        num_cols = len(fields)

        # Row 1: Fill instructions (hints for each field)
        for col_idx, (field_name, desc, hint) in enumerate(fields, 1):
            hint_text = hint if hint else (f'[{desc}]' if desc else '')
            cell_hint = ws.cell(row=1, column=col_idx, value=hint_text)
            cell_hint.font = hint_font
            cell_hint.fill = hint_fill
            cell_hint.alignment = hint_align
            cell_hint.border = thin_border

        # Row 2: Field names (headers recognized by import)
        for col_idx, (field_name, desc, hint) in enumerate(fields, 1):
            cell_fname = ws.cell(row=2, column=col_idx, value=field_name)
            cell_fname.font = header_font
            cell_fname.fill = header_fill
            cell_fname.alignment = header_align
            cell_fname.border = thin_border

        # Row 3: Empty example row (data starts here)
        for col_idx in range(1, num_cols + 1):
            cell_data = ws.cell(row=3, column=col_idx, value='')
            cell_data.font = data_font
            cell_data.border = thin_border

        # ── Column widths ──
        for col_idx, (field_name, desc, hint) in enumerate(fields, 1):
            width = max(len(field_name) * 2 + 4, 14)
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

        # ── Note at bottom ──
        note_row = 5
        note_cell = ws.cell(row=note_row, column=1, value='📌 导入步骤：从第3行开始填入数据，保存后直接上传即可（第1行说明和空白行会自动跳过）')
        note_cell.font = Font(name='微软雅黑', size=10, color='FF6600', italic=True)
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=num_cols)

        # ── Freeze ──
        ws.freeze_panes = 'A3'
        ws.auto_filter.ref = f'A2:{openpyxl.utils.get_column_letter(num_cols)}2'

        # Save to response
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        filename = f'{model_type}_import_template.xlsx'
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        wb.save(response)
        return response


importer_api_urls = [
    path('template/', DataImportTemplateView.as_view(), name='api-importer-template'),
    path('models/', DataImporterModelList.as_view(), name='api-importer-model-list'),
    path(
        'session/',
        include([
            path(
                '<int:pk>/',
                include([
                    path(
                        'accept_fields/',
                        DataImportSessionAcceptFields.as_view(),
                        name='api-import-session-accept-fields',
                    ),
                    path(
                        'accept_rows/',
                        DataImportSessionAcceptRows.as_view(),
                        name='api-import-session-accept-rows',
                    ),
                    path(
                        '',
                        DataImportSessionDetail.as_view(),
                        name='api-import-session-detail',
                    ),
                ]),
            ),
            path('', DataImportSessionList.as_view(), name='api-importer-session-list'),
        ]),
    ),
    path(
        'column-mapping/',
        include([
            path(
                '<int:pk>/',
                DataImportColumnMappingDetail.as_view(),
                name='api-importer-mapping-detail',
            ),
            path(
                '',
                DataImportColumnMappingList.as_view(),
                name='api-importer-mapping-list',
            ),
        ]),
    ),
    path(
        'row/',
        include([
            path(
                '<int:pk>/',
                DataImportRowDetail.as_view(),
                name='api-importer-row-detail',
            ),
            path('', DataImportRowList.as_view(), name='api-importer-row-list'),
        ]),
    ),
]

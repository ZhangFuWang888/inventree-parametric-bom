"""Test parametric BOM demo functionality."""
import os, sys, io
import django

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..', 'InvenTree'))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'InvenTree.settings')
django.setup()

from django.test import RequestFactory
from django.contrib.auth.models import User
from rest_framework.test import force_authenticate
from parametric_bom.api import export_bom_csv, export_attachment_zip, export_bundle_zip
from parametric_bom.bom_expander import evaluate_part
from part.models import Part

admin = User.objects.get(pk=1)
product = Part.objects.get(pk=2779)
factory = RequestFactory()

passed = 0
failed = 0

def check(name, condition, detail=''):
    global passed, failed
    if condition:
        passed += 1
        print(f'  PASS [{name}]')
    else:
        failed += 1
        print(f'  FAIL [{name}] {detail}')

# ── Test 1: BOM Evaluate ──
print('=== Test 1: BOM Evaluate ===')
result = evaluate_part(product, {'L': 3000, 'W': 500, 'P': 120, 'V': 15})
bt = result['bom_tree']
children = bt.get('children', [])
check('6 BOM items', len(children) == 6, f'Got {len(children)}')
for c in children:
    name = c.get('variant_name') or c.get('actual_part_name') or c.get('part_name', '?')
    qty = c.get('calculated_quantity', 1)
    mode = c.get('mode', 'static')
    print(f'    {name:20s} x{qty:3.0f} [{mode}]')
# Check specific quantities
qty_map = {}
for c in children:
    name = c.get('variant_name') or c.get('actual_part_name') or c.get('part_name', '')
    qty = c.get('calculated_quantity', 1)
    for key in ['机架体', '侧板', '辊筒', '支腿', '驱动', '连接件']:
        if key in name:
            qty_map[key] = qty
check('机架 x1', qty_map.get('机架体') == 1)
check('侧板 x2', qty_map.get('侧板') == 2)
check('辊筒 x26', qty_map.get('辊筒') == 26)
check('支腿 x6', qty_map.get('支腿') == 6)
check('驱动 x1', qty_map.get('驱动') == 1)
check('连接件 x52', qty_map.get('连接件') == 52)

# ── Test 2: BOM XLSX Export ──
print('\n=== Test 2: BOM XLSX Export ===')
req = factory.get('/api/parametric-bom/export/bom-csv/?part_id=2779')
force_authenticate(req, admin)
resp = export_bom_csv(req)
check('Status 200', resp.status_code == 200, str(resp.status_code))
ct = resp.get('Content-Type', '')
check('Content-Type xlsx', 'spreadsheetml' in ct, ct)
body = resp.getvalue() if hasattr(resp, 'getvalue') else b''

import openpyxl
wb = openpyxl.load_workbook(io.BytesIO(body))
ws = wb.active
check('Has header', ws.cell(1, 1).value == '序号', str(ws.cell(1, 1).value))
check('Has IPN col', ws.cell(1, 2).value == '图号/代号')
check('Has Name col', ws.cell(1, 3).value == '名称')
check('Has Qty col', ws.cell(1, 4).value == '数量')
check('Has data rows', ws.max_row > 2, f'{ws.max_row} rows')
print(f'    Header: {[ws.cell(1,c).value for c in range(1, ws.max_column+1)]}')
print(f'    Row 2: {[str(ws.cell(2,c).value) for c in range(1, ws.max_column+1)]}')
wb.close()

# ── Test 3: Attachment ZIP ──
print('\n=== Test 3: Attachment ZIP ===')
req = factory.get('/api/parametric-bom/export/attachment-zip/?part_id=2779')
force_authenticate(req, admin)
resp = export_attachment_zip(req)
# 404 is expected (no attachments), or if it happens to find some, 200 is also fine
check('Status ok (404=no attachments)', resp.status_code in (200, 404), str(resp.status_code))

# ── Test 4: Bundle ZIP ──
print('\n=== Test 4: Bundle ZIP ===')
req = factory.get('/api/parametric-bom/export/bundle-zip/?part_id=2779')
force_authenticate(req, admin)
resp = export_bundle_zip(req)
check('Status 200', resp.status_code == 200)
body = resp.getvalue() if hasattr(resp, 'getvalue') else b''
import zipfile
zf = zipfile.ZipFile(io.BytesIO(body))
names = zf.namelist()
check('ZIP has BOM.xlsx', 'BOM清单.xlsx' in names, str(names))
check('ZIP has attachments/', any(n.startswith('附件/') for n in names) or len(names) == 1,
      f'Files: {names}')
# Verify XLSX inside ZIP
wb2 = openpyxl.load_workbook(io.BytesIO(zf.read('BOM清单.xlsx')))
ws2 = wb2.active
check('XLSX has data', ws2.max_row > 2, f'{ws2.max_row} rows')
wb2.close()
zf.close()

# ── Test 5: Dynamic BOM with different params ──
print('\n=== Test 5: Dynamic BOM (L=2000, P=100, W=600) ===')
result2 = evaluate_part(product, {'L': 2000, 'W': 600, 'P': 100, 'V': 12})
bt2 = result2['bom_tree']
qty_map2 = {}
for c in bt2.get('children', []):
    name = c.get('variant_name') or c.get('actual_part_name') or c.get('part_name', '')
    qty = c.get('calculated_quantity', 1)
    for key in ['机架体', '侧板', '辊筒', '支腿', '驱动', '连接件']:
        if key in name:
            qty_map2[key] = qty
    # Also print
    vname = c.get('variant_name', '')
    pname = c.get('actual_part_name') or c.get('part_name', '?')
    print(f'    {vname or pname:25s} x{qty:3.0f}')
check('辊筒 x21 (L=2000,P=100)', qty_map2.get('辊筒') == 21,
      f'Got {qty_map2.get(chr(34)+chr(34))}' if False else '')
check('支腿 x4 (L=2000)', qty_map2.get('支腿') == 4)
check('连接件 x42', qty_map2.get('连接件') == 42)

print(f'\n{"="*40}')
print(f'Results: {passed} passed, {failed} failed')
print(f'{"="*40}')
sys.exit(0 if failed == 0 else 1)
